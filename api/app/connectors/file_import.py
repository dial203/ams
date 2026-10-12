"""Generic CSV / XLSX import connector with header-similarity mapping suggestions."""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from difflib import SequenceMatcher
from typing import Any, ClassVar

from app.connectors.base import ConnectorError, MappingSuggestion, Table
from app.forms.schema import index_schema

MAX_ROWS = 50_000
ATHLETE_HEADERS = (
    "athlete_code",
    "code",
    "athlete_id",
    "external_id",
    "player_id",
    "participant_id",
    "athlete",
    "athlete_name",
    "player",
    "player_name",
    "participant",
    "name",
    "full_name",
    "email",
    "athlete_email",
)
DATE_HEADERS = (
    "local_date",
    "date",
    "session_date",
    "test_date",
    "day",
    "observed_at",
    "datetime",
    "date_time",
    "timestamp",
    "time_stamp",
    "start_time",
)
SIMILARITY_THRESHOLD = 0.75


def normalize_header(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def _cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0) and value.tzinfo is None:
            return value.date().isoformat()
        return value.isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def _unique_columns(raw: list[str]) -> list[str]:
    out: list[str] = []
    for i, name in enumerate(raw):
        base = name.strip() or f"column_{i + 1}"
        candidate, n = base, 1
        while candidate in out:
            n += 1
            candidate = f"{base}_{n}"
        out.append(candidate)
    return out


class FileImportConnector:
    key: ClassVar[str] = "file"
    name: ClassVar[str] = "CSV / Excel upload"

    def read(self, filename: str, data: bytes) -> Table:
        lower = filename.lower()
        if lower.endswith(".xlsx"):
            return self._read_xlsx(data)
        if lower.endswith((".csv", ".txt", ".tsv")):
            return self._read_csv(data)
        raise ConnectorError("unsupported file type; upload .csv or .xlsx")

    # ------------------------------------------------------------------ readers
    def _read_csv(self, data: bytes) -> Table:
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:  # pragma: no cover - latin-1 always decodes
            raise ConnectorError("cannot decode file")
        if not text.strip():
            raise ConnectorError("file is empty")
        sample = text[:8192]
        try:
            dialect: Any = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = [r for r in reader if any(c.strip() for c in r)]
        if not rows:
            raise ConnectorError("file is empty")
        columns = _unique_columns(rows[0])
        if len(rows) - 1 > MAX_ROWS:
            raise ConnectorError(f"too many rows (max {MAX_ROWS})")
        out = [{c: (r[i].strip() if i < len(r) else "") for i, c in enumerate(columns)} for r in rows[1:]]
        return Table(columns=columns, rows=out)

    def _read_xlsx(self, data: bytes) -> Table:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for corrupt files
            raise ConnectorError("cannot read the Excel file") from exc
        try:
            ws = wb.worksheets[0]
            rows = [[_cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        finally:
            wb.close()
        rows = [r for r in rows if any(c for c in r)]
        if not rows:
            raise ConnectorError("file is empty")
        width = max(len(r) for r in rows)
        header = rows[0] + [""] * (width - len(rows[0]))
        while header and not header[-1]:
            header.pop()
        columns = _unique_columns(header)
        if len(rows) - 1 > MAX_ROWS:
            raise ConnectorError(f"too many rows (max {MAX_ROWS})")
        out = [{c: (r[i] if i < len(r) else "") for i, c in enumerate(columns)} for r in rows[1:]]
        return Table(columns=columns, rows=out)

    # ------------------------------------------------------------------ suggestions
    def suggest(self, table: Table, schema: dict[str, Any]) -> MappingSuggestion:
        athlete_col = _pick(table.columns, ATHLETE_HEADERS)
        date_col = _pick([c for c in table.columns if c != athlete_col], DATE_HEADERS)
        reserved = {athlete_col, date_col}
        fields = [(f.key, str(f.spec.get("label") or f.key)) for f in index_schema(schema).input_fields]
        scored: list[tuple[float, str, str]] = []
        for col in table.columns:
            if col in reserved:
                continue
            norm = normalize_header(col)
            for key, label in fields:
                scored.append((_similarity(norm, key, normalize_header(label)), col, key))
        mapping: dict[str, str] = {}
        used: set[str] = set()
        for score, col, key in sorted(scored, key=lambda t: (-t[0], table.columns.index(t[1]))):
            if score < SIMILARITY_THRESHOLD or col in mapping or key in used:
                continue
            mapping[col] = key
            used.add(key)
        ordered = {c: mapping[c] for c in table.columns if c in mapping}
        return MappingSuggestion(mapping=ordered, athlete_column=athlete_col, date_column=date_col)

    def fetch(self, start: dt.date, end: dt.date) -> Table:  # pragma: no cover
        raise NotImplementedError("file connectors are upload-only")


def _similarity(norm_col: str, key: str, norm_label: str) -> float:
    if not norm_col:
        return 0.0
    if norm_col in (key, norm_label):
        return 1.0
    stripped = re.sub(r"_(cm|kg|m|min|s|ms|w|au|pct|percent)$", "", norm_col)
    if stripped and stripped in (key, norm_label, re.sub(r"_(cm|kg|m|min|s|ms|w|au)$", "", key)):
        return 0.95
    abbrev = norm_col.rstrip("s")
    if len(abbrev) >= 4 and (key.startswith(abbrev) or norm_label.startswith(abbrev)):
        return 0.8  # "Accels" -> accelerations, "Decel" -> decelerations
    return max(SequenceMatcher(None, norm_col, key).ratio(), SequenceMatcher(None, norm_col, norm_label).ratio())


def _pick(columns: list[str], preferred: tuple[str, ...]) -> str | None:
    norms = {normalize_header(c): c for c in reversed(columns)}
    for name in preferred:
        if name in norms:
            return norms[name]
    return None
